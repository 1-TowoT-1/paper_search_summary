from __future__ import annotations

import hashlib
import json
import mimetypes
import re
from io import BytesIO
from dataclasses import dataclass
from pathlib import Path
from uuid import UUID, uuid4

import fitz
from fastapi import UploadFile
from sqlalchemy.orm import Session

from app.config import settings
from app.llm.llm_client import LLMClient
from app.models.db import Paper, PaperChunk, PaperFile, Project, ProjectPaper
from app.models.schemas import ManualMaterialRequest, PublicationStatus, UploadMaterialResponse, Visibility
from app.services.embedding_service import EmbeddingService
from app.services.vector_store import VectorStore


class MaterialIngestionError(RuntimeError):
    pass


class OCRExtractor:
    async def extract_pdf(self, path: Path) -> str:
        if not settings.pdf_ocr_enabled:
            raise MaterialIngestionError("OCR is disabled. Set PDF_OCR_ENABLED=true to enable it.")
        provider = settings.pdf_ocr_provider.lower().strip()
        if provider != "tesseract":
            raise MaterialIngestionError(f"Unsupported OCR provider: {settings.pdf_ocr_provider}")
        return self._extract_pdf_with_tesseract(path)

    def _extract_pdf_with_tesseract(self, path: Path) -> str:
        try:
            from PIL import Image
            import pytesseract
        except ImportError as exc:
            raise MaterialIngestionError(
                "Tesseract OCR dependencies are not installed. Run `pip install pytesseract pillow` "
                "and install the Tesseract executable."
            ) from exc

        if settings.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd

        texts: list[str] = []
        for page_number, image_bytes in _render_pdf_pages(path, max_pages=settings.pdf_ocr_max_pages):
            image = Image.open(BytesIO(image_bytes))
            text = pytesseract.image_to_string(image, lang=settings.pdf_ocr_languages).strip()
            if text:
                texts.append(f"[page {page_number}]\n{text}")
        if not texts:
            raise MaterialIngestionError("OCR did not extract readable text from PDF")
        return "\n\n".join(texts)


class MultimodalExtractor:
    def __init__(self) -> None:
        self.llm = LLMClient()

    async def extract_pdf(self, path: Path) -> str:
        if not settings.pdf_multimodal_enabled:
            raise MaterialIngestionError("Multimodal extraction is disabled. Set PDF_MULTIMODAL_ENABLED=true to enable it.")
        transport = settings.pdf_multimodal_image_transport.lower().strip()
        images = [image for _, image in _render_pdf_pages(path, max_pages=settings.pdf_multimodal_max_pages)]
        if not images:
            raise MaterialIngestionError("PDF pages could not be rendered for multimodal extraction")

        if transport == "url":
            try:
                image_urls = _write_pdf_page_urls(
                    images=images,
                    namespace=f"user_material_{uuid4().hex}",
                )
                return (await self.llm.extract_text_from_image_urls(image_urls=image_urls, filename=path.name)).strip()
            except Exception:
                return (await self.llm.extract_text_from_images(images=images, filename=path.name)).strip()

        return (await self.llm.extract_text_from_images(images=images, filename=path.name)).strip()


@dataclass
class ExtractedText:
    text: str
    method: str
    pages: list[tuple[int | None, str]]


@dataclass
class MaterialMetadata:
    title: str = ""
    abstract: str = ""
    authors: list[dict] | None = None
    keywords: list[str] | None = None
    language: str = ""
    confidence: str = "low"
    method: str = ""


def _render_pdf_pages(path: Path, max_pages: int) -> list[tuple[int, bytes]]:
    images: list[tuple[int, bytes]] = []
    matrix = fitz.Matrix(settings.pdf_render_zoom, settings.pdf_render_zoom)
    with fitz.open(path) as document:
        page_count = min(max_pages, document.page_count)
        for page_index in range(page_count):
            page = document.load_page(page_index)
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            images.append((page_index + 1, pixmap.tobytes("png")))
    return images


def _render_pdf_page_urls(path: Path, max_pages: int, namespace: str) -> list[str]:
    images = _render_pdf_pages(path, max_pages=max_pages)
    return _write_pdf_page_urls(images=images, namespace=namespace)


def _write_pdf_page_urls(images: list[tuple[int, bytes]], namespace: str) -> list[str]:
    urls: list[str] = []
    output_dir = Path(settings.generated_image_dir) / namespace
    output_dir.mkdir(parents=True, exist_ok=True)
    base_url = settings.generated_image_base_url.rstrip("/")
    for page_number, image_bytes in images:
        filename = f"page_{page_number}.png"
        (output_dir / filename).write_bytes(image_bytes)
        urls.append(f"{base_url}/{namespace}/{filename}")
    return urls


class UserMaterialImporter:
    def __init__(self) -> None:
        self.embeddings = EmbeddingService()
        self._vector_store: VectorStore | None = None
        self.llm = LLMClient()
        self.ocr = OCRExtractor()
        self.multimodal = MultimodalExtractor()

    @property
    def vector_store(self) -> VectorStore:
        if self._vector_store is None:
            self._vector_store = VectorStore()
        return self._vector_store

    async def upload_material(
        self,
        db: Session,
        file: UploadFile,
        user_id: UUID,
        project_id: UUID | None,
        title: str | None,
        abstract: str | None,
        authors: list[dict],
        publication_status: PublicationStatus,
        visibility: Visibility,
    ) -> UploadMaterialResponse:
        stats: dict[str, int | str | list[str]] = {"chunks_created": 0, "chunks_vectorized": 0, "errors": []}
        original_filename = self._safe_filename(file.filename or "uploaded_material")
        file_bytes = await file.read()
        self._validate_file_size(file_bytes)
        sha256 = hashlib.sha256(file_bytes).hexdigest()

        existing_file = (
            db.query(PaperFile)
            .filter(PaperFile.user_id == user_id, PaperFile.sha256 == sha256)
            .order_by(PaperFile.created_at.desc())
            .first()
        )
        if existing_file:
            if project_id:
                self._link_project(db, project_id=project_id, paper_id=existing_file.paper_id, user_id=user_id)
                db.commit()
            return UploadMaterialResponse(
                paper_id=existing_file.paper_id,
                status="duplicate",
                message="同一用户已上传过相同文件，已返回已有资料。",
                stats={"duplicate": 1, "sha256": sha256},
            )

        stored_path = self._store_file(user_id=user_id, sha256=sha256, filename=original_filename, content=file_bytes)
        mime_type = file.content_type or mimetypes.guess_type(original_filename)[0] or "application/octet-stream"
        extracted = await self._extract_text(stored_path, mime_type=mime_type)
        if not extracted.text.strip():
            raise MaterialIngestionError("文件未能解析出可信文本，已拒绝入库。")

        multimodal_metadata = await self._extract_multimodal_metadata(stored_path) if stored_path.suffix.lower() == ".pdf" else None
        metadata_authors = multimodal_metadata.authors if multimodal_metadata and multimodal_metadata.authors else []
        resolved_authors = authors or metadata_authors
        resolved_title = self._resolve_title(
            user_title=title,
            filename=original_filename,
            text=extracted.text,
            model_title=multimodal_metadata.title if multimodal_metadata else "",
        )
        model_abstract = None
        resolved_abstract = (abstract or "").strip()
        if not resolved_abstract:
            model_abstract = await self._resolve_abstract_from_text(
                extracted.text,
                model_abstract=multimodal_metadata.abstract if multimodal_metadata else "",
            )
            resolved_abstract = model_abstract
        if not resolved_abstract.strip():
            raise MaterialIngestionError("未提供摘要，且系统无法生成导入摘要。")

        paper = Paper(
            id=uuid4(),
            title=resolved_title,
            authors=resolved_authors,
            abstract=resolved_abstract,
            doi=None,
            source="user_upload",
            source_id=self._user_source_id(user_id=user_id, sha256=sha256),
            published_date=None,
            pdf_url=None,
            citation_count=0,
            metadata_json={
                "user_material": True,
                "model_abstract": model_abstract,
                "abstract_overridden_by_user": bool(abstract),
                "title_overridden_by_user": bool(title),
                "multimodal_metadata": self._metadata_to_dict(multimodal_metadata),
            },
            source_type="user_upload",
            publication_status=publication_status.value,
            visibility=visibility.value,
            owner_user_id=user_id,
            original_filename=original_filename,
            file_path=str(stored_path),
            file_mime_type=mime_type,
            file_sha256=sha256,
            ingestion_status="processing",
            analysis_status="pending",
            metadata_confidence=self._metadata_confidence(title=title, abstract=abstract, metadata=multimodal_metadata),
            text_extraction_method=extracted.method,
        )
        db.add(paper)
        db.flush()

        db.add(
            PaperFile(
                id=uuid4(),
                paper_id=paper.id,
                user_id=user_id,
                original_filename=original_filename,
                stored_path=str(stored_path),
                mime_type=mime_type,
                file_size_bytes=len(file_bytes),
                sha256=sha256,
                upload_status="completed",
            )
        )

        chunks = self._chunk_extracted_text(extracted)
        if not chunks:
            raise MaterialIngestionError("文件解析成功，但正文分块为空。")

        for chunk_index, (chunk_text, page_start, page_end) in enumerate(chunks):
            chunk_id = uuid4()
            db.add(
                PaperChunk(
                    id=chunk_id,
                    paper_id=paper.id,
                    chunk_index=chunk_index,
                    text=chunk_text,
                    page_start=page_start,
                    page_end=page_end,
                    chunk_type="body",
                    metadata_json={"source_type": "user_upload", "filename": original_filename},
                )
            )
            stats["chunks_created"] = int(stats["chunks_created"]) + 1

            vector = await self.embeddings.embed_text(chunk_text)
            await self.vector_store.upsert_chunk(
                paper_id=paper.id,
                chunk_id=str(chunk_id),
                vector=vector,
                text=chunk_text,
                metadata={
                    "title": resolved_title,
                    "source": "user_upload",
                    "source_id": paper.source_id,
                    "locator": f"chunk:{chunk_index}",
                    "user_id": str(user_id),
                    "project_id": str(project_id) if project_id else "",
                    "source_type": "user_upload",
                    "publication_status": publication_status.value,
                    "visibility": visibility.value,
                    "chunk_type": "body",
                },
            )
            stats["chunks_vectorized"] = int(stats["chunks_vectorized"]) + 1

        abstract_vector = await self.embeddings.embed_text(f"{resolved_title}\n\n{resolved_abstract}")
        await self.vector_store.upsert_paper_abstract(
            paper_id=paper.id,
            vector=abstract_vector,
            metadata={
                "title": resolved_title,
                "source": "user_upload",
                "source_id": paper.source_id,
                "published_date": "",
                "user_id": str(user_id),
                "project_id": str(project_id) if project_id else "",
                "source_type": "user_upload",
                "publication_status": publication_status.value,
                "visibility": visibility.value,
            },
        )

        if project_id:
            self._link_project(db, project_id=project_id, paper_id=paper.id, user_id=user_id)

        paper.ingestion_status = "completed"
        db.commit()
        db.refresh(paper)
        stats["sha256"] = sha256
        stats["text_extraction_method"] = extracted.method
        if multimodal_metadata and multimodal_metadata.method:
            stats["metadata_extraction_method"] = multimodal_metadata.method
        return UploadMaterialResponse(
            paper_id=paper.id,
            status="completed",
            message="用户资料已上传、解析并写入知识库。",
            stats=stats,
        )

    async def create_manual_material(
        self,
        db: Session,
        payload: ManualMaterialRequest,
        user_id: UUID,
    ) -> UploadMaterialResponse:
        stats: dict[str, int | str | list[str]] = {"chunks_created": 0, "chunks_vectorized": 0, "errors": []}
        content = (payload.content or payload.abstract).strip()
        source_hash = hashlib.sha256(f"{user_id}:{payload.title}:{payload.abstract}".encode("utf-8")).hexdigest()

        existing = (
            db.query(Paper)
            .filter(Paper.owner_user_id == user_id, Paper.source == "manual", Paper.source_id == source_hash)
            .first()
        )
        if existing:
            if payload.project_id:
                self._link_project(db, project_id=payload.project_id, paper_id=existing.id, user_id=user_id)
                db.commit()
            return UploadMaterialResponse(
                paper_id=existing.id,
                status="duplicate",
                message="同一用户已创建过相同手动资料。",
                stats={"duplicate": 1},
            )

        paper = Paper(
            id=uuid4(),
            title=payload.title.strip(),
            authors=payload.authors,
            abstract=payload.abstract.strip(),
            doi=None,
            source="manual",
            source_id=source_hash,
            published_date=None,
            pdf_url=None,
            citation_count=0,
            metadata_json=payload.metadata,
            source_type="manual",
            publication_status=payload.publication_status.value,
            visibility=payload.visibility.value,
            owner_user_id=user_id,
            ingestion_status="processing",
            analysis_status="pending",
            metadata_confidence="high",
            text_extraction_method="manual",
        )
        db.add(paper)
        db.flush()

        for chunk_index, chunk_text in enumerate(self._chunk_text(content)):
            chunk_id = uuid4()
            db.add(
                PaperChunk(
                    id=chunk_id,
                    paper_id=paper.id,
                    chunk_index=chunk_index,
                    text=chunk_text,
                    chunk_type="body",
                    metadata_json={"source_type": "manual"},
                )
            )
            stats["chunks_created"] = int(stats["chunks_created"]) + 1
            vector = await self.embeddings.embed_text(chunk_text)
            await self.vector_store.upsert_chunk(
                paper_id=paper.id,
                chunk_id=str(chunk_id),
                vector=vector,
                text=chunk_text,
                metadata={
                    "title": paper.title,
                    "source": "manual",
                    "source_id": paper.source_id,
                    "locator": f"chunk:{chunk_index}",
                    "user_id": str(user_id),
                    "project_id": str(payload.project_id) if payload.project_id else "",
                    "source_type": "manual",
                    "publication_status": payload.publication_status.value,
                    "visibility": payload.visibility.value,
                    "chunk_type": "body",
                },
            )
            stats["chunks_vectorized"] = int(stats["chunks_vectorized"]) + 1

        abstract_vector = await self.embeddings.embed_text(f"{paper.title}\n\n{paper.abstract}")
        await self.vector_store.upsert_paper_abstract(
            paper_id=paper.id,
            vector=abstract_vector,
            metadata={
                "title": paper.title,
                "source": "manual",
                "source_id": paper.source_id,
                "published_date": "",
                "user_id": str(user_id),
                "project_id": str(payload.project_id) if payload.project_id else "",
                "source_type": "manual",
                "publication_status": payload.publication_status.value,
                "visibility": payload.visibility.value,
            },
        )

        if payload.project_id:
            self._link_project(db, project_id=payload.project_id, paper_id=paper.id, user_id=user_id)

        paper.ingestion_status = "completed"
        db.commit()
        db.refresh(paper)
        return UploadMaterialResponse(
            paper_id=paper.id,
            status="completed",
            message="手动资料已写入个人资料库。",
            stats=stats,
        )

    async def _extract_text(self, path: Path, mime_type: str) -> ExtractedText:
        suffix = path.suffix.lower()
        if mime_type == "application/pdf" or suffix == ".pdf":
            return await self._extract_pdf(path)
        if suffix == ".docx":
            return self._extract_docx(path)
        if suffix in {".md", ".markdown"}:
            return self._extract_markdown(path)
        raise MaterialIngestionError("第一版仅支持 PDF、DOCX 和 Markdown 文件。")

    async def _extract_pdf(self, path: Path) -> ExtractedText:
        pages: list[tuple[int | None, str]] = []
        with fitz.open(path) as document:
            for page_index, page in enumerate(document, start=1):
                text = re.sub(r"\s+", " ", page.get_text("text")).strip()
                if text:
                    pages.append((page_index, text))
        text = "\n\n".join(page_text for _, page_text in pages).strip()

        errors: list[str] = []
        try:
            multimodal_text = (await self.multimodal.extract_pdf(path)).strip()
            if multimodal_text:
                combined = self._combine_extracted_text(
                    primary_text=multimodal_text,
                    secondary_text=text,
                    marker="NATIVE_TEXT_FALLBACK",
                )
                method = "multimodal+native" if text else "multimodal"
                return ExtractedText(text=combined, method=method, pages=[(None, combined)])
        except Exception as exc:
            errors.append(str(exc))

        if len(text) >= settings.pdf_min_native_text_chars:
            return ExtractedText(text=text, method="native", pages=pages)

        try:
            ocr_text = (await self.ocr.extract_pdf(path)).strip()
            if ocr_text:
                combined = self._combine_extracted_text(
                    primary_text=ocr_text,
                    secondary_text=text,
                    marker="NATIVE_TEXT_FALLBACK",
                )
                method = "ocr+native" if text else "ocr"
                return ExtractedText(text=combined, method=method, pages=[(None, combined)])
        except Exception as exc:
            errors.append(str(exc))

        if text:
            return ExtractedText(text=text, method="native", pages=pages)
        raise MaterialIngestionError("PDF 多模态、原生解析和 OCR 均失败：" + " | ".join(errors))

    def _extract_docx(self, path: Path) -> ExtractedText:
        try:
            from docx import Document
        except ImportError as exc:
            raise MaterialIngestionError("python-docx is not installed. Run `pip install -r requirements.txt`.") from exc

        document = Document(str(path))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        if not text:
            raise MaterialIngestionError("DOCX 文件未解析出正文文本。")
        return ExtractedText(text=text, method="docx", pages=[(None, text)])

    def _extract_markdown(self, path: Path) -> ExtractedText:
        try:
            text = path.read_text(encoding="utf-8").strip()
        except UnicodeDecodeError:
            text = path.read_text(encoding="gbk").strip()
        if not text:
            raise MaterialIngestionError("Markdown 文件为空。")
        return ExtractedText(text=text, method="markdown", pages=[(None, text)])

    async def _extract_multimodal_metadata(self, path: Path) -> MaterialMetadata | None:
        if not settings.pdf_multimodal_enabled:
            return None
        try:
            images = _render_pdf_pages(path, max_pages=settings.pdf_multimodal_max_pages)
            if settings.pdf_multimodal_image_transport.lower().strip() == "url":
                try:
                    image_urls = _write_pdf_page_urls(
                        images=images,
                        namespace=f"user_material_metadata_{uuid4().hex}",
                    )
                    raw_metadata = await self.llm.extract_metadata_from_image_urls(image_urls=image_urls, filename=path.name)
                except Exception:
                    raw_metadata = await self.llm.extract_metadata_from_images(
                        images=[image for _, image in images],
                        filename=path.name,
                    )
            else:
                raw_metadata = await self.llm.extract_metadata_from_images(
                    images=[image for _, image in images],
                    filename=path.name,
                )
        except Exception:
            return None

        authors = raw_metadata.get("authors")
        if not isinstance(authors, list):
            authors = []
        keywords = raw_metadata.get("keywords")
        if not isinstance(keywords, list):
            keywords = []
        confidence = str(raw_metadata.get("confidence") or "low").strip().lower()
        if confidence not in {"high", "medium", "low"}:
            confidence = "low"
        return MaterialMetadata(
            title=str(raw_metadata.get("title") or "").strip(),
            abstract=str(raw_metadata.get("abstract") or "").strip(),
            authors=[item for item in authors if isinstance(item, dict)],
            keywords=[str(item).strip() for item in keywords if str(item).strip()],
            language=str(raw_metadata.get("language") or "").strip(),
            confidence=confidence,
            method="multimodal",
        )

    def _metadata_to_dict(self, metadata: MaterialMetadata | None) -> dict:
        if metadata is None:
            return {}
        return {
            "title": metadata.title,
            "abstract": metadata.abstract,
            "authors": metadata.authors or [],
            "keywords": metadata.keywords or [],
            "language": metadata.language,
            "confidence": metadata.confidence,
            "method": metadata.method,
        }

    def _metadata_confidence(
        self,
        title: str | None,
        abstract: str | None,
        metadata: MaterialMetadata | None,
    ) -> str:
        if title or abstract:
            return "high"
        if metadata and metadata.confidence:
            return metadata.confidence
        return "low"

    async def _resolve_abstract_from_text(self, text: str, model_abstract: str = "") -> str:
        extracted_abstract = self._extract_declared_abstract(text)
        if extracted_abstract:
            return extracted_abstract
        if model_abstract.strip():
            return model_abstract.strip()
        try:
            generated_abstract = await self._generate_ingestion_abstract(text)
        except Exception:
            generated_abstract = ""
        if generated_abstract:
            return generated_abstract
        return self._fallback_extractive_abstract(text)

    async def _generate_ingestion_abstract(self, text: str) -> str:
        abstract = await self.llm.generate_ingestion_abstract(text)
        return abstract.strip()

    def _extract_declared_abstract(self, text: str) -> str:
        normalized = re.sub(r"\r\n?", "\n", text).strip()
        patterns = [
            r"(?is)(?:^|\n)\s*abstract\s*[:.\-]?\s*(.{120,2500}?)(?=\n\s*(?:keywords?|index terms|introduction|1\s*\.?\s*introduction)\b)",
            r"(?s)(?:^|\n)\s*摘要\s*[:：]?\s*(.{80,1800}?)(?=\n\s*(?:关键词|关键字|引言|导言|Abstract)\b)",
        ]
        for pattern in patterns:
            match = re.search(pattern, normalized)
            if match:
                return re.sub(r"\s+", " ", match.group(1)).strip()[:1500]

        compact = re.sub(r"\s+", " ", normalized)
        compact_patterns = [
            r"(?is)\babstract\s*[:.\-]?\s*(.{120,2500}?)(?=\b(?:keywords?|index terms|introduction|1\s*\.?\s*introduction|review article|original article)\b)",
            r"(?is)\brezumat\s*[:.\-]?\s*(.{120,2500}?)(?=\b(?:cuvinte\s+cheie|keywords?|review article|original article|chirurgia|introduction)\b)",
            r"(?s)摘要\s*[:：]?\s*(.{80,1800}?)(?=(?:关键词|关键字|引言|导言|Abstract))",
        ]
        for pattern in compact_patterns:
            match = re.search(pattern, compact)
            if match:
                return match.group(1).strip()[:1500]
        return ""

    def _fallback_extractive_abstract(self, text: str) -> str:
        cleaned = re.sub(r"\s+", " ", text).strip()
        if len(cleaned) < 80:
            return ""
        return cleaned[:1200]

    def _combine_extracted_text(self, primary_text: str, secondary_text: str, marker: str) -> str:
        primary_text = primary_text.strip()
        secondary_text = secondary_text.strip()
        if not primary_text:
            return secondary_text
        if not secondary_text:
            return primary_text
        if secondary_text in primary_text:
            return primary_text
        return f"{primary_text}\n\n[{marker}]\n\n{secondary_text}"

    def _chunk_extracted_text(self, extracted: ExtractedText) -> list[tuple[str, int | None, int | None]]:
        chunks: list[tuple[str, int | None, int | None]] = []
        if len(extracted.pages) > 1:
            for page_number, page_text in extracted.pages:
                for chunk in self._chunk_text(page_text):
                    chunks.append((chunk, page_number, page_number))
            return chunks
        return [(chunk, None, None) for chunk in self._chunk_text(extracted.text)]

    def _chunk_text(self, text: str) -> list[str]:
        normalized = re.sub(r"\s+", " ", text).strip()
        if not normalized:
            return []
        chunk_size = settings.pdf_chunk_size
        overlap = min(settings.pdf_chunk_overlap, chunk_size // 2)
        chunks: list[str] = []
        start = 0
        while start < len(normalized):
            end = min(start + chunk_size, len(normalized))
            chunk = normalized[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end >= len(normalized):
                break
            start = max(0, end - overlap)
        return chunks

    def _store_file(self, user_id: UUID, sha256: str, filename: str, content: bytes) -> Path:
        root = Path(settings.upload_dir)
        target_dir = root / str(user_id) / sha256[:2]
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / f"{sha256}_{filename}"
        with target.open("wb") as output:
            output.write(content)
        return target

    def _validate_file_size(self, content: bytes) -> None:
        max_bytes = settings.upload_file_max_size_mb * 1024 * 1024
        if len(content) > max_bytes:
            raise MaterialIngestionError(f"文件超过 {settings.upload_file_max_size_mb}MB 限制。")

    def _safe_filename(self, filename: str) -> str:
        cleaned = re.sub(r"[^\w.\-\u4e00-\u9fff]+", "_", filename).strip("._")
        return cleaned or "uploaded_material"

    def _resolve_title(self, user_title: str | None, filename: str, text: str, model_title: str = "") -> str:
        if user_title and user_title.strip():
            return user_title.strip()
        if model_title.strip():
            return model_title.strip()[:240]
        for line in text.splitlines():
            cleaned = line.strip().lstrip("#").strip()
            if 5 <= len(cleaned) <= 240:
                return cleaned
        return Path(filename).stem[:240] or "Untitled material"

    def _user_source_id(self, user_id: UUID, sha256: str) -> str:
        return hashlib.sha256(f"{user_id}:{sha256}".encode("utf-8")).hexdigest()

    def _link_project(self, db: Session, project_id: UUID, paper_id: UUID, user_id: UUID) -> None:
        project = db.query(Project).filter(Project.id == project_id, Project.user_id == user_id).first()
        if not project:
            raise MaterialIngestionError("Project not found or not owned by current user")
        existing = db.get(ProjectPaper, {"project_id": project_id, "paper_id": paper_id})
        if existing:
            return
        db.add(ProjectPaper(project_id=project_id, paper_id=paper_id))
        db.flush()


def parse_authors_json(value: str | None) -> list[dict]:
    if not value:
        return []
    try:
        parsed = json.loads(value)
    except json.JSONDecodeError as exc:
        raise MaterialIngestionError("authors_json must be valid JSON") from exc
    if not isinstance(parsed, list):
        raise MaterialIngestionError("authors must be a JSON list")
    return [item for item in parsed if isinstance(item, dict)]
